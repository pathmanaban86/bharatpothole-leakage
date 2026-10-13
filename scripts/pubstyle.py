#!/usr/bin/env python3
"""
Shared figure style, colours and notation (imported by 07_make_figures.py).
Running this file directly rebuilds the PRE-REVISION figures and tables (B-clean = 337-image list);
the published versions come from 07_make_figures.py and the manuscript tables.

Publication figures and tables - BharatPotHole leakage paper (Springer Nature style)
====================================================================================
Reads only the result files already on Drive; trains and predicts nothing.

INPUT  (MyDrive/POTHOLE/...)
  experiments/leakage_paper/results/T0_audit_numbers.json
  experiments/leakage_paper/results/T1_T3_main_results.csv
  experiments/leakage_paper/results/T4_dose_response.csv
  experiments/leakage_paper/results/T6_positives_only.csv
  dataset_assessment/per_image.csv, dataset_assessment/leakage_summary.csv
  BharatPotHole_v2/manifest_v2.csv, BharatPotHole_v2/build_summary.json

OUTPUT (MyDrive/POTHOLE/experiments/leakage_paper/publication/)
  figures/Fig1..Fig5 .pdf (vector, fonts embedded) .tiff (600 dpi, LZW) .png (preview)
  tables/Tables.docx (Word, Springer layout)  tables/Tables.tex (sn-jnl)  tables/Table*.csv
  figure_captions.md  (captions that use the same notation as the figures and tables)

STYLE
  Springer widths: single column 84 mm, double column 174 mm. Arial-compatible sans serif
  (Liberation Sans in Colab), 7-8 pt. White background, no grids, no titles inside figures
  (titles go in captions), bold lowercase panel labels. Okabe-Ito colours, checked for
  colour-vision deficiency: red-orange = leakage, blue = leakage-free, green = second model.
  Tables: caption above, three horizontal rules, no vertical lines, footnotes a, b, c.

RUN IN COLAB
  from google.colab import drive; drive.mount('/content/drive')
  !python "/content/drive/MyDrive/POTHOLE/scripts/make_publication_assets.py"
"""
import json
import os
import re
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd

# ----------------------------------------------------------------------------- paths
POTHOLE = Path(os.environ.get('POTHOLE_DRIVE', '/content/drive/MyDrive/POTHOLE'))
RES     = POTHOLE / 'experiments' / 'leakage_paper' / 'results'
PUB     = POTHOLE / 'experiments' / 'leakage_paper' / 'publication'
FIG, TAB = PUB / 'figures', PUB / 'tables'

# ----------------------------------------------------------------------------- notation
# One place for every symbol and name, used by figures, tables and captions alike.
AP50   = r'AP$_{50}$'
AP5095 = r'AP$_{50:95}$'
DAP50  = r'$\Delta$AP$_{50}$'
DH     = r'$\mathit{d}_{\mathrm{H}}$'
RHO    = r'$\mathit{\rho}$'
SET_NAME = {'original test (A)': 'A-test', 'leakage-free subset (B)': 'B-clean',
            'leaked subset': 'Leaked', 'v2 test (C)': 'C-test',
            'copies of train frames (100% leaked)': 'Train copies',
            'RDD2022-India (external)': 'RDD-India'}
MODEL_NAME = {'yolo26n': 'YOLO26n', 'yolo26s': 'YOLO26s'}
BIN_NAME = {'pHash 0-4': '0–4', 'pHash 5-8': '5–8', 'pHash 9-12': '9–12', 'pHash >12': '>12'}

# Okabe-Ito subset (validated: CVD ΔE ≥ 11, normal-vision ΔE ≥ 25, contrast ≥ 3:1 on white)
C_LEAK, C_CLEAN, C_ALT = '#D55E00', '#0072B2', '#009E73'
C_NEUTRAL, C_LIGHT, C_INK = '#7F7F7F', '#C8C8C8', '#222222'

MM = 1 / 25.4
W1, W15, W2 = 84 * MM, 129 * MM, 174 * MM


# ============================================================================ figure style
def setup_style():
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.font_manager as fm
    import matplotlib.pyplot as plt
    lib = Path('/usr/share/fonts/truetype/liberation')
    if 'Arial' not in {f.name for f in fm.fontManager.ttflist} and not lib.exists():
        subprocess.run('apt-get -qq install -y fonts-liberation > /dev/null 2>&1', shell=True)
    for f in lib.glob('LiberationSans*.ttf'):        # register regular, bold and italic weights
        fm.fontManager.addfont(str(f))
    plt.rcParams.update({
        'font.family': 'sans-serif',
        'font.sans-serif': ['Arial', 'Liberation Sans', 'Helvetica', 'DejaVu Sans'],
        'mathtext.default': 'regular',              # subscripts in the same font as the text
        'mathtext.fontset': 'custom', 'mathtext.rm': 'sans', 'mathtext.it': 'sans:italic',
        'mathtext.bf': 'sans:bold', 'mathtext.cal': 'sans', 'mathtext.sf': 'sans',
        'font.size': 7.5, 'axes.labelsize': 7.5, 'xtick.labelsize': 7, 'ytick.labelsize': 7,
        'legend.fontsize': 7, 'axes.linewidth': 0.6, 'lines.linewidth': 1.2,
        'xtick.major.width': 0.6, 'ytick.major.width': 0.6,
        'xtick.major.size': 2.5, 'ytick.major.size': 2.5,
        'axes.spines.top': False, 'axes.spines.right': False, 'axes.grid': False,
        'axes.facecolor': 'white', 'figure.facecolor': 'white', 'savefig.facecolor': 'white',
        'legend.frameon': False, 'legend.handlelength': 1.6,
        'pdf.fonttype': 42, 'ps.fonttype': 42,      # embedded, editable TrueType fonts
        'figure.constrained_layout.use': True,
    })
    return plt


def save(fig, name):
    fig.savefig(FIG / f'{name}.pdf')
    fig.savefig(FIG / f'{name}.tiff', dpi=600, pil_kwargs={'compression': 'tiff_lzw'})
    fig.savefig(FIG / f'{name}.png', dpi=300)
    print('  saved', name)


def panel(ax, letter):
    ax.annotate(letter, xy=(0, 1), xycoords='axes fraction', xytext=(-30, 6), textcoords='offset points',
                fontweight='bold', fontsize=9, ha='left', va='bottom')


# ============================================================================ load inputs
def load():
    need = [RES / 'T0_audit_numbers.json', RES / 'T1_T3_main_results.csv',
            RES / 'T4_dose_response.csv', RES / 'T6_positives_only.csv']
    miss = [str(p) for p in need if not p.exists()]
    if miss:
        sys.exit('Missing result files:\n  ' + '\n  '.join(miss))
    d = {'audit': json.load(open(need[0])), 'main': pd.read_csv(need[1]),
         'dose': pd.read_csv(need[2]), 'pos': pd.read_csv(need[3])}
    d['main']['set'] = d['main'].test_set.map(SET_NAME)
    d['main']['protocol'] = d['main'].trained_on.str[0]
    d['pos'] = d['pos'].set_index('set')
    d['pi'] = pd.read_csv(POTHOLE / 'dataset_assessment' / 'per_image.csv')
    d['man'] = pd.read_csv(POTHOLE / 'BharatPotHole_v2' / 'manifest_v2.csv')
    ls = POTHOLE / 'dataset_assessment' / 'leakage_summary.csv'
    bs = POTHOLE / 'BharatPotHole_v2' / 'build_summary.json'
    d['leak'] = pd.read_csv(ls).iloc[0] if ls.exists() else None
    d['build'] = json.load(open(bs))['checks'] if bs.exists() else {}
    return d


def row(main, model, protocol, test_set):
    r = main[(main.model == model) & (main.protocol == protocol) & (main.set == test_set)]
    return r.iloc[0] if len(r) else None


# ============================================================================ figures
def fig1_audit(plt, d):
    """Released vs unique images, split into positive and negative frames."""
    a = d['audit']
    data = [('Released', a['positive_images_released'], a['negative_images_released']),
            ('Unique', a['positive_images_unique'], a['negative_images_unique'])]
    fig, ax = plt.subplots(figsize=(W1, 1.35))
    for i, (_, p, n) in enumerate(data):
        ax.barh(i, p, height=0.6, color=C_CLEAN, edgecolor='white', linewidth=1.5)
        ax.barh(i, n, left=p, height=0.6, color=C_LIGHT, edgecolor='white', linewidth=1.5)
        ax.text(p / 2, i, f'{p:,}', ha='center', va='center', color='white', fontsize=7)
        ax.text(p + n / 2, i, f'{n:,}', ha='center', va='center', color=C_INK, fontsize=7)
    ax.set_yticks([0, 1], [r[0] for r in data])
    ax.invert_yaxis()
    ax.set_xlabel('Images')
    ax.set_xlim(0, max(p + n for _, p, n in data) * 1.02)
    ax.spines['left'].set_visible(False)
    ax.tick_params(axis='y', length=0)
    from matplotlib.patches import Patch
    ax.legend([Patch(color=C_CLEAN), Patch(color=C_LIGHT)], ['Positive', 'Negative'],
              ncol=2, loc='lower right', bbox_to_anchor=(1.0, 1.0), borderaxespad=0)
    save(fig, 'Fig1_dataset_audit')
    plt.close(fig)


def fig2_distance(plt, d):
    """d_H from each held-out image to its nearest training image, before and after re-splitting."""
    pi, man = d['pi'], d['man']
    orig = pi.loc[pi.split.isin(['valid', 'test']), 'nn_phash_to_train'].dropna()
    v2 = man.loc[man.new_split.isin(['valid', 'test']), 'nn_phash_to_train_v2'].dropna()
    bins = np.arange(0, 33)
    fig, ax = plt.subplots(figsize=(W1, 1.9))
    for data, c, ls, lab in [(orig, C_LEAK, '-', 'Published split'),
                             (v2, C_CLEAN, (0, (3, 1.5)), 'Video-disjoint split')]:
        h, _ = np.histogram(data, bins)
        ax.step(bins[:-1], 100 * h / len(data), where='post', color=c, ls=ls, label=lab)
    ax.axvline(4.5, color=C_NEUTRAL, lw=0.8, ls=(0, (3, 2)))
    ax.set_xlabel(DH + ' to nearest training image')
    ax.set_ylabel('Held-out images (%)')
    ax.set_xlim(0, 32)
    ax.set_ylim(bottom=0)
    ax.legend(loc='upper right')
    save(fig, 'Fig2_phash_distance')
    plt.close(fig)


def fig3_inflation(plt, d):
    """a: same model on Leaked / A-test / B-clean; b: AP50 versus d_H (positive images only)."""
    p = d['pos']
    fig, (a, b) = plt.subplots(1, 2, figsize=(W2, 2.2), gridspec_kw={'width_ratios': [1, 1.25]})
    names = [('leaked subset', 'Leaked', C_LEAK), ('original test (A)', 'A-test', C_NEUTRAL),
             ('leakage-free subset (B)', 'B-clean', C_CLEAN)]
    for i, (key, lab, c) in enumerate(names):
        r = p.loc[key]
        a.bar(i, r.AP50, width=0.6, color=c)
        a.errorbar(i, r.AP50, yerr=[[r.AP50 - r.AP50_lo], [r.AP50_hi - r.AP50]],
                   color=C_INK, lw=0.8, capsize=2.5, capthick=0.8)
    a.set_xticks(range(3), [n[1] for n in names])
    a.tick_params(axis='x', length=0)
    a.set_ylabel(AP50)
    a.set_ylim(0, 1.0)
    diff = p.loc['DIFF original test (A) minus leakage-free subset (B)', 'AP50']
    top = p.loc['original test (A)', 'AP50_hi'] + 0.06
    a.plot([1, 1, 2, 2], [top - 0.02, top, top, top - 0.02], color=C_INK, lw=0.7)
    a.text(1.5, top + 0.015, f'{DAP50} = {diff:.2f}', ha='center', va='bottom', fontsize=7)
    panel(a, 'a')

    keys = list(BIN_NAME)
    y = np.array([p.loc[k, 'AP50'] for k in keys])
    lo = np.array([p.loc[k, 'AP50_lo'] for k in keys])
    hi = np.array([p.loc[k, 'AP50_hi'] for k in keys])
    x = np.arange(len(keys))
    b.errorbar(x, y, yerr=[y - lo, hi - y], color=C_INK, lw=1.0, marker='o', ms=4,
               mfc=C_INK, capsize=2.5, capthick=0.8)
    b.set_xticks(x, [BIN_NAME[k] for k in keys])
    b.set_xlabel(DH + ' to nearest training image')
    b.set_ylabel(AP50)
    b.set_ylim(0, 1.0)
    b.set_xlim(-0.4, len(keys) - 0.6)
    panel(b, 'b')
    save(fig, 'Fig3_same_model_inflation')
    plt.close(fig)


def fig4_dose(plt, d):
    """AP50 of video-disjoint models as rho of positive test images are replaced by train copies."""
    dose, main = d['dose'], d['main']
    fig, ax = plt.subplots(figsize=(W1, 2.1))
    style = {'yolo26n': (C_CLEAN, 'o'), 'yolo26s': (C_ALT, 's')}
    for m, g in dose.groupby('model'):
        g = g.sort_values('leaked_fraction_of_positives')
        x = 100 * g.leaked_fraction_of_positives
        c, mk = style.get(m, (C_NEUTRAL, '^'))
        ax.fill_between(x, g.AP50_mean - g.AP50_sd, g.AP50_mean + g.AP50_sd, color=c, alpha=0.18, lw=0)
        ax.plot(x, g.AP50_mean, color=c, marker=mk, ms=3.5, label=MODEL_NAME.get(m, m))
    ref = row(main, 'yolo26n', 'A', 'A-test')
    if ref is not None:
        ax.axhline(ref.AP50, color=C_LEAK, lw=0.8, ls=(0, (3, 2)))
        ax.text(2, ref.AP50 + 0.012, 'YOLO26n, Protocol A, A-test', color=C_LEAK, fontsize=6.5, va='bottom')
    ax.set_xlabel(RHO + ' (%)')
    ax.set_ylabel(AP50)
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 0.75)
    ax.legend(loc='lower right')
    save(fig, 'Fig4_dose_response')
    plt.close(fig)


def fig5_external(plt, d):
    """Internal test versus external RDD-India, per model and training protocol."""
    main = d['main']
    series = [('yolo26n', 'A', 'A-test', C_LEAK, 'o', 'YOLO26n, Protocol A'),
              ('yolo26n', 'C', 'C-test', C_CLEAN, 's', 'YOLO26n, Protocol C'),
              ('yolo26s', 'C', 'C-test', C_ALT, '^', 'YOLO26s, Protocol C')]
    lstyle = {0: '-', 1: (0, (3, 1.5)), 2: (0, (1, 1.2))}
    fig, ax = plt.subplots(figsize=(W1, 2.1))
    off = {0: -0.03, 1: 0.0, 2: 0.03}
    for i, (m, prot, internal, c, mk, lab) in enumerate(series):
        r_in, r_ex = row(main, m, prot, internal), row(main, m, prot, 'RDD-India')
        if r_in is None or r_ex is None:
            continue
        xs = np.array([0, 1]) + off[i]
        ys = np.array([r_in.AP50, r_ex.AP50])
        err = np.array([[r_in.AP50 - r_in.AP50_lo, r_ex.AP50 - r_ex.AP50_lo],
                        [r_in.AP50_hi - r_in.AP50, r_ex.AP50_hi - r_ex.AP50]])
        ax.plot(xs, ys, color=c, lw=1.0, ls=lstyle[i])
        ax.errorbar(xs, ys, yerr=np.nan_to_num(err), fmt=mk, color=c, ms=4, lw=0.8,
                    capsize=2, capthick=0.8, label=lab)
    ax.set_xticks([0, 1], ['Internal test', 'RDD-India'])
    ax.tick_params(axis='x', length=0)
    ax.set_xlim(-0.35, 1.35)
    ax.set_ylim(0, 0.75)
    ax.set_ylabel(AP50)
    ax.legend(loc='upper right')
    save(fig, 'Fig5_external_validity')
    plt.close(fig)


# ============================================================================ table markup
# Cell markup shared by Word and LaTeX:  _{x} subscript   ^{x} superscript   {i:x} italic
TOKEN = re.compile(r'(_\{[^}]*\}|\^\{[^}]*\}|\{i:[^}]*\})')
LATEX_MAP = {'\\': r'\textbackslash{}', '&': r'\&', '%': r'\%', '#': r'\#', '$': r'\$',
             '_': r'\_', '–': '--', '—': '---', '−': '$-$', '±': r'$\pm$', '×': r'$\times$',
             '≤': r'$\leq$', '≥': r'$\geq$', '>': r'$>$', '<': r'$<$', 'ρ': r'$\rho$', 'Δ': r'$\Delta$'}


def latex(text):
    out = []
    for part in TOKEN.split(str(text)):
        if not part:
            continue
        if part.startswith('_{'):
            out.append(r'$_{\mathrm{' + part[2:-1] + '}}$')
        elif part.startswith('^{'):
            out.append(r'$^{\mathrm{' + part[2:-1] + '}}$')
        elif part.startswith('{i:'):
            inner = part[3:-1]
            out.append(LATEX_MAP.get(inner, r'\textit{' + inner + '}'))
        else:
            out.append(''.join(LATEX_MAP.get(ch, ch) for ch in part))
    return ''.join(out)


def plain(text):
    return TOKEN.sub(lambda m: m.group(0)[2:-1] if not m.group(0).startswith('{i:') else m.group(0)[3:-1],
                     str(text))


# number formatting (one rule everywhere)
def f3(v):
    return '–' if pd.isna(v) else f'{v:.3f}'


def fci(v, lo, hi):
    if pd.isna(v):
        return '–'
    return f'{v:.3f}' if pd.isna(lo) else f'{v:.3f} [{lo:.3f}, {hi:.3f}]'


def fsigned(v, lo=np.nan, hi=np.nan):
    if pd.isna(v):
        return '–'
    s = f'{v:+.3f}'.replace('-', '−')
    return s if pd.isna(lo) else f'{s} [{lo:.3f}, {hi:.3f}]'.replace('[-', '[−').replace(', -', ', −')


def fint(v):
    return '–' if pd.isna(v) else f'{int(v):,}'


def fpct(v):
    return '–' if pd.isna(v) else f'{v:.1f}'


# ============================================================================ tables
def build_tables(d):
    T = []
    a, main, p, dose = d['audit'], d['main'], d['pos'], d['dose']

    # Table 1 - notation
    T.append(dict(title='Notation and abbreviations used throughout',
                  cols=['Symbol', 'Definition'],
                  rows=[
                      ['AP_{50}', 'Average precision at IoU ≥ 0.50, all-point interpolation'],
                      ['AP_{50:95}', 'Mean average precision over IoU thresholds 0.50 to 0.95 in steps of 0.05'],
                      ['ΔAP_{50}', 'Paired difference in AP_{50} between two test subsets for the same model'],
                      ['{i:R}', 'Recall at the confidence threshold that maximises F1 (IoU ≥ 0.50)'],
                      ['{i:d}_{H}', 'Hamming distance between 64-bit perceptual hashes (pHash) of two images'],
                      ['{i:ρ}', 'Share of positive test images replaced by augmented copies of training frames'],
                      ['{i:N}_{img}, {i:N}_{pos}, {i:N}_{gt}', 'Number of images, positive images and ground-truth boxes'],
                      ['Protocol A', 'Training and evaluation on the published frame-level split'],
                      ['Protocol C', 'Training and evaluation on the deduplicated, video-disjoint split'],
                      ['A-test', 'Published test set (Protocol A)'],
                      ['Leaked', 'A-test images with a near-duplicate ({i:d}_{H} ≤ 4 and SSIM confirmation) or shared source frame in training'],
                      ['B-clean', 'A-test images with no near-duplicate or shared source frame in training'],
                      ['C-test', 'Test set of the video-disjoint split'],
                      ['Train copies', 'Augmented copies of training frames of the video-disjoint split (fully leaked)'],
                      ['RDD-India', 'Pothole (D40) images of RDD2022, India subset (external test)'],
                  ], notes=[], align='ll', latex_align=r'lp{0.72\textwidth}'))

    # Table 2 - audit
    sizes = a.get('v2_split_sizes', {})
    T.append(dict(title='Composition of the released BharatPotHole dataset before and after deduplication',
                  cols=['Quantity', 'Released', 'Unique'],
                  rows=[['Images', fint(a['images_released']), fint(a['unique_frames'])],
                        ['Positive images', fint(a['positive_images_released']), fint(a['positive_images_unique'])],
                        ['Negative images', fint(a['negative_images_released']), fint(a['negative_images_unique']) + '^{a}'],
                        ['Annotated boxes', fint(a['boxes_released']), fint(a['boxes_unique_clean']) + '^{b}'],
                        ['Exact duplicates (MD5)', fint(a['exact_md5_duplicates_removed']), '0'],
                        ['Augmentation copies', fint(a['augmentation_copies_removed']), '0']],
                  notes=[('a', 'Includes images whose only boxes were malformed or degenerate and were removed during label cleaning'),
                         ('b', 'After removal of duplicate, malformed and degenerate (< 2 px) boxes'),
                         ('', f"The unique images were split video-disjointly into {fint(sizes.get('train'))} training, "
                              f"{fint(sizes.get('valid'))} validation and {fint(sizes.get('test'))} test images")],
                  align='lrr'))

    # Table 3 - leakage of the published vs video-disjoint split
    lk, bd, man = d['leak'], d['build'], d['man']
    rows = [('group', 'Published split')]
    for s, lab in [('valid', 'Validation'), ('test', 'Test')]:
        rows.append([lab,
                     fint(lk.get(f'{s}_imgs') if lk is not None else np.nan),
                     fpct(bd.get(f'{s}_pct_within_4_of_train_original')),
                     fpct(lk.get(f'{s}_sharing_source_frame_with_train_%') if lk is not None else np.nan),
                     fpct(lk.get(f'{s}_sharing_video_prefix_with_train_%') if lk is not None else np.nan),
                     f3(bd.get(f'{s}_median_nn_phash_to_train_original')).rstrip('0').rstrip('.')])
    rows.append(('group', 'Video-disjoint split'))
    for s, lab in [('valid', 'Validation'), ('test', 'Test')]:
        rows.append([lab, fint((man.new_split == s).sum()),
                     fpct(bd.get(f'{s}_pct_within_4_of_train_v2')), '0.0', '0.0',
                     f3(bd.get(f'{s}_median_nn_phash_to_train_v2')).rstrip('0').rstrip('.')])
    T.append(dict(title='Train–test overlap in the published split and in the video-disjoint split',
                  cols=['Subset', '{i:N}_{img}', '{i:d}_{H} ≤ 4 (%)^{a}', 'Shared source frame (%)',
                        'Shared video (%)', 'Median {i:d}_{H}^{b}'],
                  rows=rows,
                  notes=[('a', 'Share of images with a training image at {i:d}_{H} ≤ 4'),
                         ('b', 'Median distance to the nearest training image')],
                  align='lrrrrr'))

    # Table 4 - all protocols and test sets (all images)
    order = [('A', 'yolo26n', ['A-test', 'B-clean', 'Leaked', 'RDD-India']),
             ('C', 'yolo26n', ['C-test', 'Train copies', 'RDD-India']),
             ('C', 'yolo26s', ['C-test', 'Train copies', 'RDD-India'])]
    rows = []
    for prot, m, sets in order:
        rows.append(('group', f'{MODEL_NAME[m]}, Protocol {prot}'))
        for s in sets:
            r = row(main, m, prot, s)
            if r is None:
                continue
            name = s + ('^{a}' if s == 'Train copies' else '^{b}' if s == 'RDD-India' else '')
            rows.append([name, fint(r.n_images), fint(r.n_pos_images), fint(r.n_gt),
                         fci(r.AP50, r.AP50_lo, r.AP50_hi), fci(r.AP50_95, r.AP50_95_lo, r.AP50_95_hi), f3(r.R)])
    T.append(dict(title='Detection performance by training protocol and test set (all images)',
                  cols=['Test set', '{i:N}_{img}', '{i:N}_{pos}', '{i:N}_{gt}',
                        'AP_{50} [95% CI]', 'AP_{50:95} [95% CI]', '{i:R}'],
                  rows=rows,
                  notes=[('', '95% confidence intervals from 1,000 bootstrap resamples of videos '
                              '(of images for RDD-India). All models: input size 640 px, confidence threshold 0.001'),
                         ('a', 'No confidence interval computed; set used for the substitution experiment'),
                         ('b', 'Contains positive images only')],
                  align='lrrrrrr'))

    # Table 5 - same-model inflation, positive images only
    rows = [('group', 'By leakage status')]
    diffs = {'original test (A)': 'DIFF original test (A) minus leakage-free subset (B)',
             'leaked subset': 'DIFF leaked subset minus leakage-free subset (B)'}
    for key, lab in [('leaked subset', 'Leaked'), ('original test (A)', 'A-test'), ('leakage-free subset (B)', 'B-clean')]:
        r = p.loc[key]
        dk = diffs.get(key)
        dv = fsigned(p.loc[dk, 'AP50'], p.loc[dk, 'AP50_lo'], p.loc[dk, 'AP50_hi']) if dk in p.index else 'Reference'
        rows.append([lab, fint(r.n_images), fint(r.n_gt), fci(r.AP50, r.AP50_lo, r.AP50_hi), f3(r.AP50_95), f3(r.R), dv])
    rows.append(('group', 'By {i:d}_{H} to the nearest training image'))
    for key, lab in BIN_NAME.items():
        if key in p.index:
            r = p.loc[key]
            rows.append([lab + ('^{a}' if lab == '0–4' else ''), fint(r.n_images), fint(r.n_gt),
                         fci(r.AP50, r.AP50_lo, r.AP50_hi), f3(r.AP50_95), f3(r.R), '–'])
    T.append(dict(title='Leakage inflation for the same detector (YOLO26n, Protocol A), positive test images only',
                  cols=['Subset', '{i:N}_{pos}', '{i:N}_{gt}', 'AP_{50} [95% CI]', 'AP_{50:95}', '{i:R}',
                        'ΔAP_{50} vs B-clean [95% CI]'],
                  rows=rows,
                  notes=[('', 'Confidence intervals from 1,000 paired bootstrap resamples of videos'),
                         ('a', 'Same images as the Leaked subset')],
                  align='lrrrrrr'))

    # Table 6 - dose-response
    piv = {}
    for m, g in dose.groupby('model'):
        for _, r in g.iterrows():
            piv.setdefault(r.leaked_fraction_of_positives, {})[m] = r
    models = sorted(dose.model.unique())

    def msd(mean, sd, reps):
        return f'{mean:.3f}' if reps == 1 else f'{mean:.3f} ± {sd:.3f}'
    rows = []
    for rho in sorted(piv):
        cells = [f'{100 * rho:.0f}']
        for m in models:
            r = piv[rho].get(m)
            cells += ['–', '–'] if r is None else [msd(r.AP50_mean, r.AP50_sd, r.reps),
                                                   msd(r.AP50_95_mean, r.AP50_95_sd, r.reps)]
        rows.append(cells)
    cols = ['{i:ρ} (%)'] + [c for m in models for c in (f'{MODEL_NAME.get(m, m)} AP_{{50}}',
                                                         f'{MODEL_NAME.get(m, m)} AP_{{50:95}}')]
    T.append(dict(title='Effect of controlled leakage on video-disjoint models (C-test with substituted positive images)',
                  cols=cols, rows=rows,
                  notes=[('', 'Values are mean ± standard deviation over 20 random substitutions; '
                              '{i:ρ} = 0 and 100% are single deterministic evaluations')],
                  align='l' + 'r' * (len(cols) - 1)))
    return T


# ----------------------------------------------------------------------------- writers
def write_csv(T):
    for i, t in enumerate(T, 1):
        body = [r for r in t['rows'] if not isinstance(r, tuple)]
        pd.DataFrame([[plain(c) for c in r] for r in body],
                     columns=[plain(c) for c in t['cols']]).to_csv(TAB / f'Table{i}.csv', index=False)


def write_latex(T):
    out = ['% Springer Nature LaTeX template (sn-jnl.cls) - tables generated automatically', '']
    for i, t in enumerate(T, 1):
        n = len(t['cols'])
        out += [r'\begin{table}[h]', r'\caption{' + latex(t['title']) + r'}\label{tab' + str(i) + '}%',
                r'\begin{tabular}{@{}' + t.get('latex_align', t['align']) + '@{}}', r'\toprule',
                ' & '.join(latex(c) for c in t['cols']) + r' \\', r'\midrule']
        first = True
        for r in t['rows']:
            if isinstance(r, tuple):
                if not first:
                    out.append(r'\midrule')
                out.append(r'\multicolumn{' + str(n) + r'}{@{}l}{\textit{' + latex(r[1]) + r'}} \\')
            else:
                out.append(' & '.join(latex(c) for c in r) + r' \\')
            first = False
        out += [r'\botrule', r'\end{tabular}']
        for mark, txt in t['notes']:
            out.append(r'\footnotetext{' + (latex('^{' + mark + '}') + ' ' if mark else '') + latex(txt) + '}')
        out += [r'\end{table}', '']
    (TAB / 'Tables.tex').write_text('\n'.join(out), encoding='utf-8')


def write_docx(T):
    try:
        import docx  # noqa: F401
    except ImportError:
        subprocess.run([sys.executable, '-m', 'pip', 'install', '-q', 'python-docx'], check=True)
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    def runs(par, text, size, bold=False, italic=False):
        for part in TOKEN.split(str(text)):
            if not part:
                continue
            if part.startswith('_{'):
                r = par.add_run(part[2:-1]); r.font.subscript = True
            elif part.startswith('^{'):
                r = par.add_run(part[2:-1]); r.font.superscript = True
            elif part.startswith('{i:'):
                r = par.add_run(part[3:-1]); r.italic = True
            else:
                r = par.add_run(part)
            r.font.size, r.font.name = Pt(size), 'Times New Roman'
            r.bold = r.bold or bold
            if italic:
                r.italic = True

    def borders(cell, top=False, bottom=False, sz=8):
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement('w:tcBorders')
        for edge, on in (('top', top), ('bottom', bottom)):     # schema order: top before bottom
            if on:
                e = OxmlElement(f'w:{edge}')
                for k, v in (('val', 'single'), ('sz', str(sz)), ('space', '0'), ('color', '000000')):
                    e.set(qn(f'w:{k}'), v)
                b.append(e)
        tcPr.append(b)

    def tight(par):
        par.paragraph_format.space_before = Pt(1)
        par.paragraph_format.space_after = Pt(1)

    from docx.shared import Cm

    def row_flag(row, tag):
        trPr = row._tr.get_or_add_trPr()
        e = OxmlElement(f'w:{tag}'); e.set(qn('w:val'), 'true'); trPr.append(e)

    def widths(t, usable=17.6):
        """Column widths (cm) from content length; headers count by their longest word."""
        w = []
        for j, c in enumerate(t['cols']):
            body = [len(plain(r[j])) for r in t['rows'] if not isinstance(r, tuple)]
            head = max(len(x) for x in plain(c).split())
            w.append(max(1.3, 0.17 * max(body + [head]) + 0.35))
        s = sum(w)
        return [x * usable / s for x in w] if s > usable else w

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(1.7)
    for i, t in enumerate(T, 1):
        cap = doc.add_paragraph()
        cap.paragraph_format.keep_with_next = True
        r = cap.add_run(f'Table {i} '); r.bold = True; r.font.size = Pt(10); r.font.name = 'Times New Roman'
        runs(cap, t['title'], 10)
        n = len(t['cols'])
        tab = doc.add_table(rows=1, cols=n)
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        tab.autofit = False
        row_flag(tab.rows[0], 'tblHeader')             # repeat header row on a new page
        for j, c in enumerate(t['cols']):
            cell = tab.rows[0].cells[j]
            par = cell.paragraphs[0]; tight(par)
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if t['align'][j] == 'l' else WD_ALIGN_PARAGRAPH.CENTER
            runs(par, c, 9)
            borders(cell, top=True, bottom=True, sz=8)
        body = t['rows']
        for k, rr in enumerate(body):
            cells = tab.add_row().cells
            last = k == len(body) - 1
            if isinstance(rr, tuple):
                merged = cells[0].merge(cells[-1])
                par = merged.paragraphs[0]; tight(par)
                runs(par, rr[1], 9, italic=True)
                borders(merged, top=k > 0, bottom=last, sz=4)
                continue
            for j, c in enumerate(rr):
                par = cells[j].paragraphs[0]; tight(par)
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT if t['align'][j] == 'l' else WD_ALIGN_PARAGRAPH.CENTER
                runs(par, c, 9)
                if last:
                    borders(cells[j], bottom=True, sz=8)
        cw = widths(t)
        for j, col in enumerate(tab.columns):
            col.width = Cm(cw[j])
        for ri, r_ in enumerate(tab.rows):
            row_flag(r_, 'cantSplit')
            for cell in r_.cells:                       # keep the whole table on one page
                for par in cell.paragraphs:
                    par.paragraph_format.keep_with_next = ri < len(tab.rows) - 1
            if len({id(c._tc) for c in r_.cells}) == n:   # skip merged group-header rows
                for j, cell in enumerate(r_.cells):
                    cell.width = Cm(cw[j])
        for mark, txt in t['notes']:
            par = doc.add_paragraph(); tight(par)
            if mark:
                runs(par, '^{' + mark + '}', 8)
            runs(par, txt, 8)
        doc.add_paragraph()
    doc.save(TAB / 'Tables.docx')


def write_captions(d):
    m = d['main']
    a_ref = row(m, 'yolo26n', 'A', 'A-test')
    txt = f"""# Figure captions (notation as in Table 1)

**Fig. 1** Composition of the released BharatPotHole dataset and of its unique images after removal of
exact duplicates and augmentation copies. Positive images contain at least one annotated pothole.

**Fig. 2** Distribution of {plain('d_{H}')} (pHash Hamming distance) from each held-out image (validation and
test) to its nearest training image under the published split and under the video-disjoint split. The
dashed line marks the near-duplicate threshold ({plain('d_{H}')} ≤ 4).

**Fig. 3** Leakage inflation for the same detector (YOLO26n, Protocol A), positive test images only.
**a** AP50 on the Leaked subset, the full A-test and the B-clean subset; ΔAP50 is the paired difference
between A-test and B-clean. **b** AP50 by {plain('d_{H}')} between each test image and its nearest training
image. Error bars: 95% confidence intervals from 1,000 bootstrap resamples of videos.

**Fig. 4** AP50 of the video-disjoint models (Protocol C) on C-test as a share ρ of positive test images is
replaced by augmented copies of training frames. Lines: mean of 20 random substitutions; bands: ±1
standard deviation. Dashed line: YOLO26n trained and tested under Protocol A
(AP50 = {a_ref.AP50:.3f}).

**Fig. 5** AP50 on the internal test set (A-test for Protocol A, C-test for Protocol C) and on the external
RDD-India set. Error bars: 95% bootstrap confidence intervals.
""" if a_ref is not None else '# Captions: Protocol A results missing\n'
    (PUB / 'figure_captions.md').write_text(txt, encoding='utf-8')


# ============================================================================ main
def main():
    FIG.mkdir(parents=True, exist_ok=True)
    TAB.mkdir(parents=True, exist_ok=True)
    d = load()
    plt = setup_style()
    print('Figures ->', FIG)
    fig1_audit(plt, d)
    fig2_distance(plt, d)
    fig3_inflation(plt, d)
    fig4_dose(plt, d)
    fig5_external(plt, d)
    T = build_tables(d)
    write_csv(T)
    write_latex(T)
    write_docx(T)
    write_captions(d)
    print('Tables  ->', TAB, '(Tables.docx, Tables.tex, Table1-6.csv)')
    print('Captions->', PUB / 'figure_captions.md')


if __name__ == '__main__':
    main()
